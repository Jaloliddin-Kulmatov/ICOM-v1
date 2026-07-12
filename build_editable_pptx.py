#!/usr/bin/env python3
"""Build an editable PPTX: clean-plate backgrounds (Russian removed, graphics
kept) + real editable Uzbek text boxes positioned on top."""
import os
import fitz
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

import translate_pdf as T

EMU_PER_DISP = 9144          # 20in slide / 2000 display units * 914400
# Rounded bold for titles/headings/body to match the deck; plain for the rest.
FONT_BOLD = "Arial Rounded MT Bold"
FONT_REG = "Arial"


def to_disp(b):
    return (b[0] / T.K, b[1] / T.K, b[2] / T.K, b[3] / T.K)


def fit_size(text, b, fontpath, max_disp, allow_wrap, pad=0.90):
    tmp = Image.new("RGB", (10, 10))
    d = ImageDraw.Draw(tmp)
    x0, y0, x1, y1 = b
    bw, bh = (x1 - x0) * pad, (y1 - y0)
    for size in range(int(max_disp * T.K), 8, -2):
        font = ImageFont.truetype(fontpath, size)
        lines = T.wrap(d, text, font, bw) if allow_wrap else [text]
        if not allow_wrap and d.textlength(text, font=font) > bw:
            continue
        asc, desc = font.getmetrics()
        lh = asc + desc
        gap = int(lh * 0.12)
        total = len(lines) * lh + (len(lines) - 1) * gap
        widest = max(d.textlength(ln, font=font) for ln in lines)
        if widest <= bw and total <= bh:
            return size
    return 12


def main():
    doc = fitz.open(T.SRC)
    os.makedirs("bg_out", exist_ok=True)
    bgs = []
    # 1) clean-plate backgrounds: remove Russian, keep graphics + pill shapes
    for i in range(doc.page_count):
        page = doc[i]
        pix = page.get_pixmap(matrix=fitz.Matrix(T.SCALE, T.SCALE))
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        for (bg, b, text, color, fp, align, mx, wr) in T.PAGES.get(i + 1, []):
            if bg == T.WHITE:
                T.cover(img, b, bg)
            else:
                T.cover_pill(img, b, bg)
        path = f"bg_out/page{i+1}.png"
        img.save(path)
        bgs.append(path)

    # 2) assemble pptx with editable text boxes
    prs = Presentation()
    prs.slide_width = Emu(int(20 * 914400))
    prs.slide_height = Emu(int(11.25 * 914400))
    blank = prs.slide_layouts[6]
    for i, bgpath in enumerate(bgs, start=1):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(bgpath, 0, 0, width=prs.slide_width, height=prs.slide_height)
        for (bg, b, text, color, fp, align, mx, wr) in T.PAGES.get(i, []):
            pt = fit_size(text, b, fp, mx, wr) * 0.5
            d = to_disp(b)
            box = slide.shapes.add_textbox(
                Emu(int(d[0] * EMU_PER_DISP)), Emu(int(d[1] * EMU_PER_DISP)),
                Emu(int((d[2] - d[0]) * EMU_PER_DISP)), Emu(int((d[3] - d[1]) * EMU_PER_DISP)))
            tf = box.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if align == "center" else PP_ALIGN.LEFT
            run = p.add_run()
            run.text = text
            f = run.font
            is_bold = (fp == T.BOLD)
            f.size = Pt(round(pt * (0.92 if is_bold else 1.0)))  # rounded bold is wider
            f.bold = is_bold
            f.name = FONT_BOLD if is_bold else FONT_REG
            f.color.rgb = RGBColor(*color)
    prs.save("PPT_uzbek_editable.pptx")
    print("saved PPT_uzbek_editable.pptx with", len(bgs), "slides")


if __name__ == "__main__":
    main()
