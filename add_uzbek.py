#!/usr/bin/env python3
"""Add an Uzbek translation line directly below each Korean text line in the deck."""
import copy
import re
from pptx import Presentation
from lxml import etree

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
def a(tag): return f"{{{A}}}{tag}"

HANGUL = re.compile(r"[가-힣]")
UZ_FONT = "S-Core Dream 3 Light"

# Korean (stripped) -> Uzbek translation
TR = {
    "한국 전통놀이": "Koreyaning an'anaviy o'yinlari",
    "안녕하세요!": "Assalomu alaykum!",
    "한국 명절&기념일 퀴즈": "Koreya bayramlari va xotira kunlari viktorinasi",
    "한국 전통 공예 체험 - 딱지": "Koreya an'anaviy hunarmandchiligi tajribasi – Ddakji",
    "한국 전통 공예 체험 - 탈": "Koreya an'anaviy hunarmandchiligi tajribasi – Niqob (Tal)",
    "활동 소감 나누기": "Faoliyat taassurotlarini o'zaro baham ko'rish",
    "퀴즈": "Viktorina",
    "한국의 설날은 음력 1월 15일이다? ---------------- ( O / X )":
        "Koreyaning Seollal (Yangi yil) bayrami oy taqvimi bo'yicha 1-oyning 15-kunida nishonlanadi. ( O / X )",
    "한국의 추수감사절 이름은?": "Koreyaning shukronalik (hosil) bayramining nomi qanday?",
    "한국은 광복절에 태극기를 건다?  ----------------- ( O / X )":
        "Koreyada Ozodlik kunida (Gwangbokjeol) davlat bayrog'i (Taegukgi) osiladi. ( O / X )",
    "한국은 어린이를 위한 날이 있다? ---------------- ( O / X )":
        "Koreyada bolalarga atalgan maxsus kun bor. ( O / X )",
    "다섯 개의 작고 동그란 돌을 던져 손으로 잡으며 노는 놀이는?":
        "Beshta kichik dumaloq toshni otib, qo'l bilan ilib o'ynaladigan o'yin nima deb ataladi?",
    "팽이치기는 팽이를 오래 돌리는 사람이 이긴다?--------- ( O / X )":
        "Pirildoq (paxmoq) o'yinida uni eng uzoq aylantirgan kishi g'olib bo'ladi. ( O / X )",
    "딱지 만들기": "Ddakji yasash",
    "2장의 딱지종이를 접는 선을 따라 접기": "Ikkita ddakji qog'ozini buklama chiziqlari bo'ylab buklang",
    "제작법 사진 출처: 안녕미술아 유튜브": "Tayyorlash rasmlari manbasi: \"Annyeong Misura\" YouTube kanali",
    "두 장의 딱지종이를 민화가 가운데 오게 겹친 후 뒤집기":
        "Ikkita ddakji qog'ozini naqsh o'rtaga keladigan qilib ustma-ust qo'yib, ag'daring",
    "숫자가 보이도록 놓은 뒤 번호 순서대로 접기":
        "Raqamlar ko'rinib turadigan qilib qo'ygach, raqamlar tartibida buklang",
    "마지막 부분(4번)은 종이 안쪽으로 집어넣기": "Oxirgi qismni (4-raqam) qog'ozning ichiga qistirib qo'ying",
    "딱지치기": "Ddakji o'yini (Ddakjichigi)",
    "가위바위보 하기": "Tosh-qaychi-qog'oz o'ynash",
    "진 사람: 딱지 내려놓기": "Yutqazgan o'yinchi: ddakjisini yerga qo'yadi",
    "이긴 사람: 내 딱지로 상대의 딱지 내려치기": "Yutgan o'yinchi: o'z ddakjisi bilan raqib ddakjisiga uradi",
    "바닥에 놓인 상대방의 딱지를 내 딱지로 쳐서 뒤집거나 선 밖으로 밀어내는 놀이":
        "Yerda turgan raqib ddakjisini o'z ddakjing bilan urib ag'darish yoki chiziqdan tashqariga chiqarish o'yini",
    "상대 딱지가 뒤집어지거나, 바람에 넘어가면 그 딱지를 획득!":
        "Raqib ddakjisi ag'darilsa yoki uchib ketib ag'darilsa, o'sha ddakji seniki bo'ladi!",
    "딱지치기 놀이": "Ddakji o'yini",
    "다같이 딱지를 쳐 봅시다!": "Keling, hammamiz birga ddakji o'ynaymiz!",
    "제한시간: 10분": "Belgilangan vaqt: 10 daqiqa",
    "정해진 시간 안에": "Belgilangan vaqt ichida",
    "상대팀의 딱지를 더 많이 가져와 봅시다!": "raqib jamoaning ddakjilaridan ko'proq yutib olishga harakat qilamiz!",
    "한 팀의 딱지는 최대 2개까지 뺏을 수 있습니다.": "Bitta jamoadan ko'pi bilan 2 tagacha ddakji yutib olish mumkin.",
    "한 팀에게 2개를 빼앗았다면 다른 팀에게 도전!": "Bir jamoadan 2 ta yutib olsangiz, boshqa jamoaga qarshi bellashing!",
    "탈 만들기": "Niqob (Tal) yasash",
    "색종이를 찢고, 오리고, 붙여서 꾸미기": "Rangli qog'ozni yirtib, qirqib va yopishtirib bezang",
    "오늘의 활동 마무리": "Bugungi mashg'ulot yakuni",
    "모두 함께 활동 소감 이야기 나누기": "Hammamiz birgalikda mashg'ulot taassurotlarini muhokama qilamiz",
    "학습 및 활동 내용 요약": "O'rganilgan va bajarilgan ishlar mazmunini umumlashtirish",
    "감사합니다!": "Rahmat!",
}

# (slide_index, shape_id) pairs already translated -> skip
SKIP = {(5, 7)}


def uz_size_cp(k_pt):
    """Uzbek subtitle size in centipoints — fixed at 28pt for all lines."""
    return 2800


def set_typeface(rpr, name):
    for tag in ("latin", "ea", "cs"):
        el = rpr.find(a(tag))
        if el is None:
            el = etree.SubElement(rpr, a(tag))
        el.set("typeface", name)


def insert_spc_before(ppr, cp):
    """Insert <a:spcBef> after lnSpc (schema order) with cp centipoints."""
    spc = etree.Element(a("spcBef"))
    pts = etree.SubElement(spc, a("spcPts"))
    pts.set("val", str(cp))
    lnspc = ppr.find(a("lnSpc"))
    if lnspc is not None:
        lnspc.addnext(spc)
    else:
        ppr.insert(0, spc)


def build_uz_paragraph(src_p, uz_text, uz_cp):
    """Clone the Korean <a:p>, strip to one run carrying the Uzbek text."""
    p = copy.deepcopy(src_p)
    runs = p.findall(a("r"))
    keep = runs[0]
    for r in runs[1:]:
        p.remove(r)
    # set text
    t = keep.find(a("t"))
    t.text = uz_text
    # style the kept run
    rpr = keep.find(a("rPr"))
    rpr.set("sz", str(uz_cp))
    rpr.set("lang", "uz-Latn-UZ")
    for attr in ("err", "spc"):
        if attr in rpr.attrib:
            del rpr.attrib[attr]
    set_typeface(rpr, UZ_FONT)
    # style endParaRPr to match, if present
    epr = p.find(a("endParaRPr"))
    if epr is not None:
        epr.set("sz", str(uz_cp))
        set_typeface(epr, UZ_FONT)
    # a little breathing room above the Uzbek line
    ppr = p.find(a("pPr"))
    if ppr is None:
        ppr = etree.Element(a("pPr"))
        p.insert(0, ppr)
    insert_spc_before(ppr, int(uz_cp * 0.18))
    return p


def main():
    prs = Presentation("presentation.pptx")
    added = 0
    missing = set()
    for si, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            if (si, shape.shape_id) in SKIP:
                continue
            tf = shape.text_frame
            # snapshot source paragraph elements first
            src_paras = list(tf.paragraphs)
            for para in src_paras:
                txt = "".join(r.text for r in para.runs)
                if not HANGUL.search(txt):
                    continue
                key = txt.strip().replace("\xa0", " ")
                uz = TR.get(key)
                if uz is None:
                    missing.add(key)
                    continue
                k_pt = para.runs[0].font.size.pt if para.runs[0].font.size else 40.0
                uz_cp = uz_size_cp(k_pt)
                new_p = build_uz_paragraph(para._p, uz, uz_cp)
                para._p.addnext(new_p)
                added += 1
    if missing:
        print("!! MISSING TRANSLATIONS:")
        for m in sorted(missing):
            print("   ", repr(m))
    print(f"Added {added} Uzbek lines.")
    prs.save("presentation_uz.pptx")
    print("Saved presentation_uz.pptx")


if __name__ == "__main__":
    main()
