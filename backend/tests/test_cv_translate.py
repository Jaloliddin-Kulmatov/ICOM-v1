"""CV translator: uploaded file or pasted text -> Korean 이력서 / English CV, and Word export."""
import io
import json
from types import SimpleNamespace

import docx
import pytest

import routes.ai as ai_routes
import routes.cv as cv_routes
from conftest import User

CV_TEXT = """Linh Tran — Marketing student at Jeonbuk National University
Email: linh@example.com
Education: Jeonbuk National University, BA Business, 2024.03 - present
Experience: Social media intern, Bean Cafe, 2025.07 - 2025.08. Wrote Instagram posts.
Skills: Canva, Excel. Languages: Vietnamese (native), Korean (TOPIK 4)."""

KOREAN = {
    "full_name": "Linh Tran", "headline": "전북대학교 경영학과 재학 중", "email": "linh@example.com",
    "education": [{"school": "전북대학교", "degree": "학사 재학 중", "field": "경영학", "start": "2024.03", "end": "현재", "details": ""}],
    "experience": [{"title": "SNS 마케팅 인턴", "company": "Bean Cafe", "start": "2025.07", "end": "2025.08", "details": "인스타그램 게시물 작성"}],
    "skills": ["Canva", "Excel"], "languages": [{"name": "베트남어", "level": "모국어"}, {"name": "한국어", "level": "TOPIK 4급"}],
    "unknown_field": "dropped",
}


@pytest.fixture
def fake_ai(monkeypatch):
    calls = []

    def create(model, messages, **kw):
        calls.append(messages)
        reply = calls_reply[0]
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=reply))])

    calls_reply = [json.dumps(KOREAN, ensure_ascii=False)]
    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    monkeypatch.setattr(ai_routes, "_get_client", lambda: client)
    cv_routes._translate_hits.clear()
    return SimpleNamespace(calls=calls, reply=calls_reply)


def test_translate_pasted_text(client, fake_ai):
    assert client.post("/api/cv/me/translate", json={"text": CV_TEXT}).status_code == 401
    u = User(client, "Linh", "cv.text@icom.test")
    r = u.post("/api/cv/me/translate", json={"source": "text", "text": CV_TEXT, "target": "ko"})
    assert r.status_code == 200, r.get_json()
    cv = r.get_json()["cv"]
    assert cv["full_name"] == "Linh Tran" and cv["education"][0]["school"] == "전북대학교"
    assert "unknown_field" not in cv                       # cleaned like the CV maker
    prompt = fake_ai.calls[0][0]["content"]
    assert "Use ONLY facts" in prompt and "Korean" in prompt
    assert CV_TEXT in fake_ai.calls[0][1]["content"]


def test_translate_uploaded_docx(client, fake_ai):
    u = User(client, "Doc", "cv.doc@icom.test")
    assert u.post("/api/cv/me/translate", json={"source": "file", "target": "ko"}).status_code == 400   # nothing uploaded
    d = docx.Document()
    for line in CV_TEXT.split("\n"):
        d.add_paragraph(line)
    buf = io.BytesIO(); d.save(buf)
    up = u.post("/api/cv/me/file", data={"file": (io.BytesIO(buf.getvalue()), "cv.docx")}, content_type="multipart/form-data")
    assert up.status_code == 201, up.get_json()
    r = u.post("/api/cv/me/translate", json={"source": "file", "target": "en"})
    assert r.status_code == 200, r.get_json()
    assert "Jeonbuk National University" in fake_ai.calls[-1][1]["content"]   # text came out of the Word file
    assert "English" in fake_ai.calls[-1][0]["content"]


def test_translate_handles_fenced_json_and_errors(client, fake_ai):
    u = User(client, "Fence", "cv.fence@icom.test")
    fake_ai.reply[0] = "Here you go:\n```json\n" + json.dumps(KOREAN, ensure_ascii=False) + "\n```"
    assert u.post("/api/cv/me/translate", json={"text": CV_TEXT, "target": "ko"}).status_code == 200
    assert u.post("/api/cv/me/translate", json={"text": "hi", "target": "ko"}).status_code == 400
    assert u.post("/api/cv/me/translate", json={"text": CV_TEXT, "target": "fr"}).status_code == 400
    fake_ai.reply[0] = "sorry, no json"
    assert u.post("/api/cv/me/translate", json={"text": CV_TEXT, "target": "ko"}).status_code == 400


def test_translate_daily_limit(client, fake_ai):
    u = User(client, "Busy", "cv.busy@icom.test")
    for _ in range(cv_routes.TRANSLATIONS_PER_DAY):
        assert u.post("/api/cv/me/translate", json={"text": CV_TEXT, "target": "ko"}).status_code == 200
    assert u.post("/api/cv/me/translate", json={"text": CV_TEXT, "target": "ko"}).status_code == 429


def test_export_word_file(client):
    u = User(client, "Export", "cv.export@icom.test")
    r = u.post("/api/cv/me/export", json={"cv": KOREAN, "lang": "ko"})
    assert r.status_code == 200 and r.data[:2] == b"PK"
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(r.data)).paragraphs)
    cells = " ".join(c.text for t in docx.Document(io.BytesIO(r.data)).tables for row in t.rows for c in row.cells)
    assert "이력서" in text and "학력" in text and "전북대학교" in cells and "SNS 마케팅 인턴" in cells
    r = u.post("/api/cv/me/export", json={"cv": KOREAN, "lang": "en"})
    assert "Curriculum Vitae" in "\n".join(p.text for p in docx.Document(io.BytesIO(r.data)).paragraphs)
