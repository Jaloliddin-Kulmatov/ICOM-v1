"""CV translator: turn a student's CV into a Korean 이력서 or an English CV.

read_cv_file()  text out of an uploaded PDF or Word file
translate_cv()  AI rewrites the CV text into ICOM's CV sections, in Korean or English,
                using only facts that are in the text
build_docx()    a clean Word file from those sections
"""
import io
import json
import re

MAX_INPUT_CHARS = 12000
LANG_NAMES = {"ko": "Korean", "en": "English"}

SHAPE = {
    "full_name": "", "headline": "", "nationality": "", "visa": "", "email": "", "phone": "", "location": "",
    "summary": "",
    "education": [{"school": "", "degree": "", "field": "", "start": "", "end": "", "details": ""}],
    "experience": [{"title": "", "company": "", "location": "", "start": "", "end": "", "details": ""}],
    "skills": [""],
    "languages": [{"name": "", "level": ""}],
    "certifications": [{"name": "", "year": ""}],
    "links": [{"label": "", "url": ""}],
}


class CVError(Exception):
    """Something the student can fix (empty file, scanned PDF, too long...)."""


def read_cv_file(content: bytes, mime: str) -> str:
    if mime == "application/pdf":
        from pypdf import PdfReader
        try:
            reader = PdfReader(io.BytesIO(content))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception as e:  # noqa: BLE001 — broken or encrypted PDFs raise many types
            raise CVError("We couldn't read that PDF. Try a Word file, or paste the text instead.") from e
    else:
        import docx
        doc = docx.Document(io.BytesIO(content))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 40:
        raise CVError("This file has almost no text we can read (it may be a scanned image). Paste the text instead.")
    return text


def _prompt(target: str) -> str:
    lang = LANG_NAMES[target]
    korean_rules = (
        "Write it as a formal Korean 이력서: short phrases ending in nouns (예: '마케팅 콘텐츠 기획 및 운영'), "
        "dates as YYYY.MM, degree names in Korean (학사, 석사, 재학 중). "
        if target == "ko" else
        "Write clear, simple professional English with short bullet-like sentences. Dates as YYYY.MM. "
    )
    return (
        f"You turn a student's CV into a clean CV written in {lang}, for applying to jobs in Korea.\n"
        "Rules:\n"
        "1. Use ONLY facts that are in the CV text. Never add, guess or improve anything: no new jobs, schools, dates, "
        "numbers, skills or achievements. If something is missing, leave that field empty.\n"
        f"2. Translate every field into {lang}. Keep names of people, companies, schools and products as written "
        "(you may add the official Korean name in parentheses only if you are certain of it).\n"
        "3. Copy email addresses, phone numbers and URLs exactly.\n"
        f"4. {korean_rules}\n"
        "5. headline: the person's current study or role in one short line. summary: 2-4 sentences based only on the CV's "
        "own summary or objective; empty if the CV has none.\n"
        "6. In 'details', put the duties and results from the CV for that item, one per line.\n"
        "Return ONLY a JSON object with exactly this shape (lists may be empty):\n"
        + json.dumps(SHAPE, ensure_ascii=False)
    )


def _parse_json(text: str) -> dict:
    text = (text or "").strip()
    try:
        return json.loads(text)
    except ValueError:
        m = re.search(r"\{.*\}", text, re.S)   # some models wrap JSON in prose or ```json fences
        if not m:
            raise CVError("The translation didn't come back in the right format. Please try again.")
        return json.loads(m.group(0))


def translate_cv(client, text: str, target: str) -> dict:
    """Return CV sections (raw; the caller cleans them with routes.cv._clean_cv)."""
    from groq_models import complete
    if target not in LANG_NAMES:
        raise CVError("Choose Korean or English.")
    text = (text or "").strip()
    if len(text) < 40:
        raise CVError("Add your CV text first (at least a few lines).")
    if len(text) > MAX_INPUT_CHARS:
        raise CVError(f"Your CV is too long to translate at once (over {MAX_INPUT_CHARS:,} characters). Shorten it a little.")
    resp = complete(client, [
        {"role": "system", "content": _prompt(target)},
        {"role": "user", "content": "CV text:\n\"\"\"\n" + text + "\n\"\"\""},
    ], temperature=0.2, max_tokens=4000)
    data = _parse_json(resp.choices[0].message.content)
    if not isinstance(data, dict):
        raise CVError("The translation didn't come back in the right format. Please try again.")
    return data


# ── Word file ────────────────────────────────────────────────────────────────

LABELS = {
    "ko": {"title": "이력서", "personal": "인적사항", "name": "성명", "nationality": "국적", "visa": "체류자격(비자)",
           "phone": "연락처", "email": "이메일", "location": "주소", "summary": "자기소개", "education": "학력",
           "experience": "경력", "skills": "기술", "languages": "어학", "certifications": "자격증", "links": "링크",
           "period": "기간", "content": "내용"},
    "en": {"title": "Curriculum Vitae", "personal": "Personal details", "name": "Name", "nationality": "Nationality",
           "visa": "Visa", "phone": "Phone", "email": "Email", "location": "Address", "summary": "Summary",
           "education": "Education", "experience": "Experience", "skills": "Skills", "languages": "Languages",
           "certifications": "Certifications", "links": "Links", "period": "Period", "content": "Details"},
}


def _period(item: dict) -> str:
    start, end = (item.get("start") or "").strip(), (item.get("end") or "").strip()
    return f"{start} ~ {end}" if start or end else ""


def build_docx(cv: dict, lang: str) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    L = LABELS["ko" if lang == "ko" else "en"]
    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "Malgun Gothic"
    style.font.size = Pt(10)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "맑은 고딕")

    def heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x00, 0x66, 0x33)

    def two_col(rows):
        t = doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        widths = (docx.shared.Cm(4), docx.shared.Cm(12.5))
        # Fixed width + explicit borders: Pages, Hangul and phone viewers ignore "Table Grid" and auto-fit.
        # Word wants tblPr children in schema order: tblStyle, tblW, ..., tblBorders, tblLayout, ..., tblLook.
        tbl_pr = t._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        tbl_w.set(qn("w:type"), "dxa"); tbl_w.set(qn("w:w"), str(sum(w.twips for w in widths)))
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4"); b.set(qn("w:color"), "BFC5C2")
            borders.append(b)
        layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed")
        tbl_w.addnext(borders); borders.addnext(layout)
        for col, w in zip(t.columns, widths):
            col.width = w
        for left, right in rows:
            cells = t.add_row().cells
            cells[0].text = left
            cells[1].text = right
            for cell, w in zip(cells, widths):
                cell.width = w   # Word reads the cell width, other apps the column width
            shade = OxmlElement("w:shd")   # light grey label column, like a printed 이력서 form
            shade.set(qn("w:val"), "clear"); shade.set(qn("w:fill"), "F2F4F3")
            cells[0]._tc.get_or_add_tcPr().append(shade)

    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(L["title"]); r.bold = True; r.font.size = Pt(20)

    personal = [(L["name"], cv.get("full_name", "")), (L["nationality"], cv.get("nationality", "")),
                (L["visa"], cv.get("visa", "")), (L["phone"], cv.get("phone", "")),
                (L["email"], cv.get("email", "")), (L["location"], cv.get("location", ""))]
    personal = [(a, b) for a, b in personal if b]
    if cv.get("headline"):
        sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.add_run(cv["headline"]).italic = True
    if personal:
        heading(L["personal"]); two_col(personal)
    if cv.get("summary"):
        heading(L["summary"]); doc.add_paragraph(cv["summary"])
    for key in ("education", "experience"):
        items = cv.get(key) or []
        if not items:
            continue
        heading(L[key])
        rows = []
        for it in items:
            if key == "education":
                head = " · ".join(x for x in (it.get("school"), it.get("degree"), it.get("field")) if x)
            else:
                head = " · ".join(x for x in (it.get("company"), it.get("title"), it.get("location")) if x)
            body = head + ("\n" + it["details"] if it.get("details") else "")
            rows.append((_period(it), body))
        two_col(rows)
    if cv.get("languages"):
        heading(L["languages"])
        two_col([(x.get("name", ""), x.get("level", "")) for x in cv["languages"]])
    if cv.get("certifications"):
        heading(L["certifications"])
        two_col([(x.get("year", ""), x.get("name", "")) for x in cv["certifications"]])
    if cv.get("skills"):
        heading(L["skills"]); doc.add_paragraph(", ".join(cv["skills"]))
    if cv.get("links"):
        heading(L["links"])
        for x in cv["links"]:
            doc.add_paragraph(f"{x.get('label', '')}: {x.get('url', '')}".strip(": "))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
