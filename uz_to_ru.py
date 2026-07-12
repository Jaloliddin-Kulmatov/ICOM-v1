#!/usr/bin/env python3
"""Replace Uzbek subtitle lines with Russian in the Korean-games deck."""
import re
from pptx import Presentation

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
HANGUL = re.compile(r"[가-힣]")

TR = {
    "Koreya an'anaviy o'yinlar": "Корейские традиционные игры",
    # 공기놀이 Gonggi
    "Gonggi o'yini": "Игра «Конги» (Gonggi)",
    "Toshlarni yerga sochish": "Рассыпать камешки по полу",
    "Toshni yuqoriga otib, yerdagi toshni olish": "Подбросить камешек вверх и взять камешек с пола",
    "Yerda tosh qolmaguncha takrorlash": "Повторять, пока на полу не останется камешков",
    "Beshta kichik dumaloq toshni havoga otib, qo'l bilan ilib o'ynaladigan o'yin":
        "Игра, в которой подбрасывают пять маленьких круглых камешков и ловят их рукой",
    "Toshlarni qo'l ustiga qo'yib, havoda ilib olish":
        "Положить камешки на тыльную сторону ладони и поймать их в воздухе",
    "Yerdagi toshni olayotganda": "Когда берёшь камешек с пола,",
    "yon-atrofdagi toshga tegsangiz — mag'lubiyat!": "если заденешь соседний камешек — проигрыш!",
    "Bir martada olinadigan tosh sonini oshirib borib":
        "Постепенно увеличивая число камешков, берущихся за один раз,",
    "beshta bo'lganda havoda ilib oling!": "когда станет пять — поймайте их в воздухе!",
    # 제기차기 Jegichagi
    "Jegichagi o'yini": "Игра «Чеги» (Jegichagi)",
    "Jegini qo'l bilan ushlash": "Взять чеги в руку",
    "Yengil qilib yuqoriga otish": "Легко подбросить вверх",
    "Yerga tushirmasdan tepishda davom etish": "Продолжать подбивать, не давая упасть на пол",
    "Jegini oyoq bilan tepib o'ynaladigan o'yin": "Игра, в которой чеги подбивают ногой",
    # 비석치기 Biseokchigi
    "Biseokchigi o'yini": "Игра «Писок-чиги» (Biseokchigi)",
    "Kaft kattaligidagi to'rtburchak toshni qo'l bilan otib yoki oyoq bilan tepib, raqib toshini urib ag'daradigan o'yin":
        "Игра, в которой квадратным камнем размером с ладонь бросают рукой или бьют ногой, сбивая камень соперника",
    # 무궁화 꽃이 피었습니다
    "Mugunghwa guli ochildi": "«Цветок мугунхва расцвёл»",
    "Quvlovchi 'Mugunghwa guli ochildi' deb baqirayotganda unga yaqinlashib, uni urib qochadigan o'yin":
        "Игра, в которой, пока водящий кричит «Цветок мугунхва расцвёл», к нему подкрадываются, задевают и убегают",
    "Quvlovchi 'Mugunghwa guli ochildi' deb baqiradi": "Водящий кричит «Цветок мугунхва расцвёл»",
    "Quvlovchi shior aytayotganda yaqinlashish": "Приближаться, пока водящий выкрикивает фразу",
    "Quvlovchini urib qochish": "Задеть водящего и убежать",
    "Quvlovchiga yetib borgan odam": "Тот, кто добрался до водящего,",
    "ushlangan odamni ozod qilib": "освобождает пойманного",
    "birga qochib ketishi mumkin!": "и они могут убежать вместе!",
    "Quvlovchi orqasiga o'girilganda": "Когда водящий оборачивается,",
    "qimirlagan odam o'yindan chiqadi!": "тот, кто пошевелился, выбывает!",
    "quvlovchiga tutilasiz!": "вас поймает водящий!",
    # 둥글게 둥글게
    "Aylana bo'lib, aylana bo'lib": "«Круг за кругом»",
    "Qo'l ushlashib aylana yasab, qo'shiq aytib aylanish":
        "Взявшись за руки, встать в круг и кружиться, распевая песню",
    "Boshlovchi aytgan sonni diqqat bilan tinglash": "Внимательно слушать число, которое называет ведущий",
    "Songa mos ravishda o'sha sonli guruh tuzish": "Собраться в группы по названному числу",
    "Bolalar bilan qo'l ushlashib, aylana yasab, musiqa ohangida aylanadigan harakatli o'yin":
        "Подвижная игра, в которой дети, взявшись за руки, образуют круг и кружатся под музыку",
    "Guruhga qo'shila olmagan yoki belgilangan songa yetkaza olmagan odam o'yindan chiqadi!":
        "Тот, кто не попал в группу или не набрал нужное число, выбывает!",
}


def set_para_text(para, text):
    runs = para.runs
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def main():
    prs = Presentation("kruz.pptx")
    done = 0
    missed = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            for para in sh.text_frame.paragraphs:
                txt = "".join(r.text for r in para.runs).strip()
                if not txt or HANGUL.search(txt):
                    continue
                if txt.isdigit():
                    continue
                if txt in TR:
                    set_para_text(para, TR[txt])
                    done += 1
                elif re.search(r"[A-Za-z]", txt):
                    missed.append(txt)
    if missed:
        print("MISSED (still Latin):")
        for m in missed:
            print("   ", repr(m))
    print(f"Replaced {done} subtitles.")
    prs.save("Korean_Traditional_Games_KR-RU.pptx")
    print("saved Korean_Traditional_Games_KR-RU.pptx")


if __name__ == "__main__":
    main()
