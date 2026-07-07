# -*- coding: utf-8 -*-
"""Add Uzbek translations below each Korean text and normalise formatting."""
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor

NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
UZ_FONT = "S-Core Dream 3 Light"     # Latin font already used for Uzbek in the deck
DARK = RGBColor(0x35, 0x35, 0x35)    # deck's body / heading colour

# key: (slide, shape_id) -> (list_of_uzbek_lines, size_pt)
ADD = {
    # slide 1 - cover title
    (1, 7):  (["Koreys an'anaviy o'yinlari"], 40),
    # slide 2 - greeting + agenda
    (2, 13): (["Assalomu alaykum!"], 40),
    (2, 23): (["Koreya bayramlari bilan tanishuv"], 24),
    (2, 24): (["Koreya xotira kunlari bilan tanishuv"], 24),
    (2, 25): (["Koreya an'anaviy hunarmandchiligi – pirildoq (top)"], 24),
    (2, 26): (["Mashg'ulot taassurotlari bilan o'rtoqlashish"], 24),
    # slide 3
    (3, 13): (["Eslab ko'raylik"], 40),
    # slide 4 - goals
    (4, 13): (["Mashg'ulot maqsadi"], 40),
    (4, 19): (["Koreyaning an'anaviy bayramlari va milliy madaniyatini tushunish"], 24),
    (4, 20): (["Amaliy mashg'ulotlar orqali madaniy almashinuv va jamoaviy muloqot tajribasini kengaytirish"], 24),
    # slide 6 - Seollal details
    (6, 21): (["Oy taqvimi bo'yicha 1-yanvar"], 26),
    (6, 22): (["Milliy taom: tteokguk (guruch cho'zma sho'rvasi)"], 26),
    (6, 23): (["An'anaviy o'yinlar: yunnori, varrak uchirish, neolttwigi va h.k."], 26),
    # slide 7 - Jeongwol Daeboreum details
    (7, 21): (["Oy taqvimi bo'yicha 15-yanvar"], 26),
    (7, 22): (["Milliy taom: besh donli guruch, yong'oq-mag'iz"], 26),
    (7, 23): (["An'anaviy o'yinlar: jwibulnori, ganggangsullae va h.k."], 26),
    # slide 8 - Chuseok details
    (8, 21): (["Oy taqvimi bo'yicha 15-avgust"], 26),
    (8, 22): (["Milliy taom: songpyeon"], 26),
    (8, 23): (["An'anaviy o'yinlar: arqon tortish, ssireum (kurash), kamon otish va h.k."], 26),
    # slide 9 - Children's day details
    (9, 21): (["Quyosh taqvimi bo'yicha 5-may"], 26),
    (9, 22): (["Oila va do'stlar bilan birga yeyiladigan taomlar"], 26),
    (9, 23): (["Oila va do'stlar bilan birga o'ynaladigan o'yinlar"], 26),
    # slide 10 - Liberation day details
    (10, 21): (["Quyosh taqvimi bo'yicha 15-avgust"], 26),
    (10, 22): (["Yaponiya mustamlakasidan ozod bo'lish va Koreya Respublikasi hukumatining tashkil etilishi"], 26),
    (10, 23): (["Milliy bayroq (Taegukgi)ni ko'tarish"], 26),
    # slide 12 - Gonggi (jacks). id32 (tip) handled separately.
    (12, 14): (["Gonggi (toshcha o'yini)"], 40),
    (12, 25): (["Beshta kichik dumaloq toshni otib, qo'l bilan ilib o'ynaladigan o'yin"], 22),
    (12, 21): (["Toshchalarni yerga sochish"], 24),
    (12, 22): (["Bitta toshni yuqoriga otib, yerdagi toshni olish"], 24),
    (12, 23): (["Yerda tosh qolmaguncha takrorlash"], 24),
    (12, 29): (["Toshlarni qo'l ustiga qo'yib, so'ng ilib olish"], 24),
    # slide 13 - Jegichagi
    (13, 14): (["Jegichagi (patcha tepish)"], 40),
    (13, 26): (["Jegini oyoq bilan tepib o'ynaladigan o'yin"], 22),
    (13, 22): (["Jegini qo'l bilan ushlash"], 24),
    (13, 23): (["Yengilgina yuqoriga otish"], 24),
    (13, 24): (["Yerga tushirmasdan doimiy tepib turish"], 24),
    # slide 14 - Biseokchigi
    (14, 14): (["Biseokchigi (tosh ag'darish o'yini)"], 40),
    (14, 16): (["Kaft kattaligidagi to'rtburchak toshni qo'l bilan otib yoki oyoq bilan tepib, raqib toshini urib ag'daradigan o'yin"], 22),
    # slide 15 & 16 - Biseokchigi title only
    (15, 14): (["Biseokchigi (tosh ag'darish o'yini)"], 40),
    (16, 14): (["Biseokchigi (tosh ag'darish o'yini)"], 40),
    # slide 17 - Ttakjichigi
    (17, 14): (["Ttakjichigi (qog'oz plastinka o'yini)"], 40),
    (17, 25): (["Yerdagi raqib ttakjisini o'z ttakjing bilan urib ag'darish yoki chiziqdan chiqarish o'yini"], 22),
    (17, 21): (["Tosh-qaychi-qog'oz o'ynash"], 24),
    (17, 22): (["Yutqazgan: ttakjisini yerga qo'yadi"], 24),
    (17, 23): (["Yutgan: o'z ttakjisi bilan raqibnikini uradi"], 24),
    (17, 29): (["Raqib ttakjisi ag'darilsa yoki shamolda uchsa – o'sha ttakji seniki!"], 22),
    # slide 18 - Paengichigi
    (18, 14): (["Paengichigi (pirildoq aylantirish)"], 40),
    (18, 17): (["Pirildoqni qamchi bilan urib aylantiriladigan o'yin"], 22),
    # slide 20 & 21 - decorating the top
    (20, 14): (["An'anaviy pirildoqni bezash"], 40),
    (21, 14): (["An'anaviy pirildoqni bezash"], 40),
    (21, 16): (["Yog'och pirildoqni flomaster va rangli qalamlar bilan bezash"], 22),
    # slide 22 - Paengichigi (repeat)
    (22, 14): (["Paengichigi (pirildoq aylantirish)"], 40),
    (22, 17): (["Pirildoqni qamchi bilan urib aylantiriladigan o'yin"], 22),
    # slide 23 - wrap up
    (23, 13): (["Bugungi mashg'ulot yakuni"], 40),
    (23, 19): (["Hammamiz birga mashg'ulot taassurotlarini o'rtoqlashamiz"], 24),
    (23, 20): (["O'rganilgan va bajarilgan ishlar mazmunini umumlashtirish"], 24),
    # slide 24 - thank you
    (24, 7): (["Rahmat!"], 40),
}

# Existing Uzbek subtitle boxes to NORMALISE (rebuild as one clean run).
FIX = {
    (5, 19):  ("Koreya bayramlari va esdalik sanalari", 26.66),
    (6, 20):  ("Seollal / Koreya yangi yili", 24.64),
    (7, 20):  ("Jeongwol Daeboreum / Buyuk to'lin oy festivali", 24.64),
    (8, 20):  ("Chuseok / Koreya minnatdorchilik kuni", 24.64),
    (9, 20):  ("Eolin-inal / Bolalar kuni", 24.64),
    (10, 20): ("Gwangbogjeol / Koreya ozodlik kuni", 24.64),
    (11, 19): ("An'anaviy koreys o'yinlari", 26.66),
    (19, 19): ("Koreya an'anaviy hunarmandchiligi", 26.66),
}

# slide-12 tip callout: Korean lines are re-spaced tight, then a compact
# 2-line Uzbek note is appended so the whole block stays inside the bubble.
TIP_KEY = (12, 32)
TIP_UZ = [
    "Toshni olayotganda boshqasiga tegmang!",
    "Sonini oshirib, beshtada hammasini iling!",
]


def q(tag):
    return f"{{{NS}}}{tag}"


def style_run(run, size_pt, bold=None):
    f = run.font
    f.name = UZ_FONT
    f.size = Pt(size_pt)
    f.color.rgb = DARK
    if bold is not None:
        f.bold = bold
    rPr = run._r.get_or_add_rPr()
    for t in ("latin", "ea", "cs"):
        e = rPr.find(q(t))
        if e is None:
            e = rPr.makeelement(q(t), {})
            rPr.append(e)
        e.set("typeface", UZ_FONT)


def set_space_before(p, pts):
    pPr = p._p.get_or_add_pPr()
    for old in pPr.findall(q("spcBef")):
        pPr.remove(old)
    spc = pPr.makeelement(q("spcBef"), {})
    spc.append(pPr.makeelement(q("spcPts"), {"val": str(int(pts * 100))}))
    pPr.insert(0, spc)


def set_line_spacing(p, pct):
    pPr = p._p.get_or_add_pPr()
    for old in pPr.findall(q("lnSpc")):
        pPr.remove(old)
    ln = pPr.makeelement(q("lnSpc"), {})
    ln.append(pPr.makeelement(q("spcPct"), {"val": str(int(pct * 1000))}))
    pPr.insert(0, ln)


def add_uzbek(shape, lines, size_pt):
    tf = shape.text_frame
    align = None
    for p in tf.paragraphs:
        if p.text.strip():
            align = p.alignment
    gap = 8 if size_pt >= 40 else 4
    for idx, line in enumerate(lines):
        p = tf.add_paragraph()
        if align is not None:
            p.alignment = align
        set_space_before(p, gap if idx == 0 else 2)
        r = p.add_run()
        r.text = line
        style_run(r, size_pt, bold=False)


def fix_subtitle(shape, text, size_pt):
    tf = shape.text_frame
    txBody = tf._txBody
    p0 = tf.paragraphs[0]
    align = p0.alignment
    for extra_p in txBody.findall(q("p"))[1:]:
        txBody.remove(extra_p)
    for r in list(p0._p.findall(q("r"))):
        p0._p.remove(r)
    if align is not None:
        p0.alignment = align
    r = p0.add_run()
    r.text = text
    style_run(r, size_pt, bold=(True if size_pt > 26 else None))


def do_tip(shape):
    tf = shape.text_frame
    # tighten the four Korean lines so the block frees space in the bubble
    for p in tf.paragraphs:
        if p.text.strip():
            set_line_spacing(p, 108)
            set_space_before(p, 0)
    for idx, line in enumerate(TIP_UZ):
        p = tf.add_paragraph()
        p.alignment = tf.paragraphs[0].alignment
        set_line_spacing(p, 100)
        set_space_before(p, 8 if idx == 0 else 2)
        r = p.add_run()
        r.text = line
        style_run(r, 18, bold=False)


def main():
    prs = Presentation("original.pptx")
    by_id = {}
    for i, slide in enumerate(prs.slides, start=1):
        for sh in slide.shapes:
            by_id[(i, sh.shape_id)] = sh

    for key, (text, size) in FIX.items():
        fix_subtitle(by_id[key], text, size)
    for key, (lines, size) in ADD.items():
        add_uzbek(by_id[key], lines, size)
    do_tip(by_id[TIP_KEY])

    prs.save("한국 전통놀이_PPT (Uzbek).pptx")
    print("saved")


if __name__ == "__main__":
    main()
